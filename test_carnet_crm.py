# Tests du 29/10/2026 : le carnet de prospects devient un mini-CRM (export
# Excel/CSV, import d'une liste avec doublons ecartes, champs telephone,
# e-mail, source et montant de la vente).
#
# Entierement mockes (ni Notion ni reseau), meme style que test_carnet_canaux.py.
#
# Lancer : python test_carnet_crm.py

import base64
import os
import csv
import io
import json
import sys
from datetime import date
from unittest.mock import patch

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook

import portal_main
from backend.services import carnet_fichiers_service as fichiers
from backend.services import notion_service as ns
from backend.services import portal_auth_service, prospects_service as carnet

portal_auth_service.SECRET_KEY = "cle-de-test-crm"
os.environ["COACH_ONBOARD_KEY"] = "cle-coach-crm"

passed = 0
failed = 0


def check(label, condition):
    global passed, failed

    if condition:
        passed += 1
        print(f"  PASS - {label}")
    else:
        failed += 1
        print(f"  FAIL - {label}")


def statut(fonction, *args, **kwargs):
    try:
        return 200, fonction(*args, **kwargs)
    except HTTPException as error:
        return error.status_code, error.detail


CLIENT = "dddddddd-dddd-dddd-dddd-dddddddddddd"
BEARER = "Bearer " + portal_auth_service.create_session_token("client@test.com", CLIENT)
DASHBOARD = {"nom": "Marie Test", "parcours": "Atelier", "identite": {"email": "client@test.com"}, "fiches": []}
AUJOURD_HUI = date.today().isoformat()


class FauxNotion:
    def __init__(self, prospects=None):
        self.page = None
        self.ecritures = 0

        if prospects is not None:
            self.page = self._page("page-carnet", json.dumps({"version": 1, "prospects": prospects}))

    def _page(self, page_id, texte):
        return {"id": page_id, "properties": {"Données (JSON)": {"type": "rich_text", "rich_text": [{"plain_text": texte}]}}}

    def query(self, data_source_id, filter_=None):
        return [self.page] if self.page and carnet.NOM_CARNET in json.dumps(filter_ or {}) else []

    def ecrire(self, *args):
        properties = args[-1]
        self.ecritures += 1
        self.page = self._page("page-carnet", "".join(m["text"]["content"] for m in properties["Données (JSON)"]["rich_text"]))
        return self.page

    def prospects(self):
        texte = self.page["properties"]["Données (JSON)"]["rich_text"][0]["plain_text"]
        return json.loads(texte)["prospects"]


def avec(faux):
    ns._cache_clear("carnet:")
    patches = [
        patch.object(ns, "_query_data_source", side_effect=faux.query),
        patch.object(ns, "_create_page", side_effect=faux.ecrire),
        patch.object(ns, "_update_page", side_effect=faux.ecrire),
        patch.object(ns, "get_client_dashboard", return_value=DASHBOARD),
    ]
    for p in patches:
        p.start()
    return patches


def fin(patches):
    for p in patches:
        p.stop()


def b64(octets):
    return base64.b64encode(octets).decode()


ANCIEN = [
    {"id": "a1", "nom": "Garage Lubin", "contact": "0690 11 22 33", "statut": "Client", "relance": None, "note": "Contrat signé",
     "cree_le": "2026-10-01", "evenements": [{"date": "2026-10-05", "statut": "Contacté"}, {"date": "2026-10-10", "statut": "RDV"}, {"date": "2026-10-20", "statut": "Client"}]},
    {"id": "a2", "nom": "Mme Jean", "contact": "mme.jean@exemple.fr", "statut": "À contacter", "relance": "2026-11-02", "note": "=HYPERLINK(\"x\")",
     "cree_le": "2026-10-02", "evenements": []},
]

print("\n=== Champs CRM ===")
p = carnet.completer(ANCIEN[0])
check("ancien contact téléphone rangé en téléphone", p["telephone"] == "0690 11 22 33" and p["email"] == "")
p = carnet.completer(ANCIEN[1])
check("ancien contact e-mail rangé en e-mail", p["email"] == "mme.jean@exemple.fr" and p["telephone"] == "")
check("montant « 1 200,50 € » lu", carnet.montant("1 200,50 €") == 1200.5)
check("montant vide = None", carnet.montant("") is None)

faux = FauxNotion([dict(x) for x in ANCIEN])
patches = avec(faux)
try:
    v = carnet.vue_client(carnet.appliquer(CLIENT, "Marie Test", DASHBOARD, "modifier", "a1", {"montant": "1500", "source": "Recommandation", "telephone": "0690 99 88 77", "email": "garage@lubin.fr"}))
    g = next(x for x in v["prospects"] if x["id"] == "a1")
    check("modifier : montant, source, téléphone, e-mail enregistrés", g["montant"] == 1500 and g["source"] == "Recommandation" and g["email"] == "garage@lubin.fr")
    check("« contact » suit le téléphone", g["contact"] == "0690 99 88 77")
    check("CA total et panier moyen", v["ventes"] == {"ca": 1500.0, "panier_moyen": 1500.0, "clients": 1})

    for label, donnees in [("e-mail invalide refusé", {"email": "pas-un-mail"}), ("montant invalide refusé", {"montant": "beaucoup"}), ("montant négatif refusé", {"montant": -5})]:
        try:
            carnet.appliquer(CLIENT, "Marie Test", DASHBOARD, "modifier", "a1", donnees)
            check(label, False)
        except carnet.CarnetInvalide:
            check(label, True)
finally:
    fin(patches)

semaine = carnet.totaux_semaine([{**carnet.completer(ANCIEN[0]), "montant": 800}], date(2026, 10, 19))
check("le montant compte la semaine où il devient client", semaine["ventes"] == 1 and semaine["ca"] == 800)
check("pas de CA la semaine d'avant", carnet.totaux_semaine([{**ANCIEN[0], "montant": 800}], date(2026, 10, 12))["ca"] == 0)

print("\n=== Export ===")
contenu = fichiers.export_csv(ANCIEN)
check("CSV : BOM pour Excel", contenu.startswith("﻿".encode("utf-8")))
lignes = list(csv.reader(io.StringIO(contenu.decode("utf-8-sig")), delimiter=";"))
check("CSV : en-têtes en français", lignes[0][:4] == ["Nom", "Téléphone", "E-mail", "Statut"])
check("CSV : dates RDV et client reprises", lignes[1][10] == "2026-10-10" and lignes[1][11] == "2026-10-20")
check("CSV : formule neutralisée", lignes[2][7].startswith("'="))
check("CSV : téléphone gardé tel quel", lignes[1][1] == "0690 11 22 33")

classeur = load_workbook(io.BytesIO(fichiers.export_xlsx(ANCIEN)))
feuille = classeur.active
check("Excel : 1 en-tête + 2 prospects", feuille.max_row == 3 and feuille["A2"].value == "Garage Lubin")
check("Excel : relance en vraie date", feuille["E3"].value.date() == date(2026, 11, 2))
check("Excel : formule neutralisée", str(feuille["H3"].value).startswith("'="))

print("\n=== Import : lecture des fichiers ===")
csv_fr = "Prénom;Nom;Téléphone;Mail;Statut;Montant\nPaul;Durand;0690 44 55 66;paul@d.fr;RDV;\n;;;;;\nLéa;Martin;;lea@m.fr;inconnu;300\n".encode("cp1252")
lu = fichiers.lire_fichier("liste.csv", b64(csv_fr))
check("CSV Windows-1252, point-virgule, prénom + nom", [l["nom"] for l in lu] == ["Paul Durand", "Léa Martin"])

google = "First Name,Last Name,E-mail 1 - Value,Phone 1 - Value\nAnne,Roy,anne@roy.fr,+590 690 12 12 12\n".encode()
check("export Google Contacts", fichiers.lire_fichier("contacts.csv", b64(google))[0]["telephone"] == "+590 690 12 12 12")

vcf = b"BEGIN:VCARD\r\nVERSION:3.0\r\nFN:Boulangerie C\xc3\xa9leste\r\nTEL;TYPE=CELL:0690 77 77 77\r\nEND:VCARD\r\nBEGIN:VCARD\r\nN:Roy;Anne;;;\r\nEMAIL:anne@roy.fr\r\nEND:VCARD\r\n"
lu = fichiers.lire_fichier("contacts.vcf", b64(vcf))
check("vCard du téléphone : 2 contacts", [l["nom"] for l in lu] == ["Boulangerie Céleste", "Anne Roy"])

classeur = Workbook()
classeur.active.append(["Entreprise", "Tel", "Relance"])
classeur.active.append(["Hôtel Bel Air", "0590 10 20 30", date(2026, 12, 1)])
sortie = io.BytesIO()
classeur.save(sortie)
lu = fichiers.lire_fichier("prospects.xlsx", b64(sortie.getvalue()))
check("Excel : colonnes reconnues", lu[0]["nom"] == "Hôtel Bel Air" and lu[0]["telephone"] == "0590 10 20 30")

for label, nom, octets in [
    ("sans colonne utile : refusé", "x.csv", b"Ville;Pays\nPAP;FR\n"),
    ("format inconnu : refusé", "x.pdf", b"%PDF"),
    ("fichier vide : refusé", "x.csv", b""),
    ("trop lourd : refusé", "x.csv", b"Nom\n" + b"a" * (fichiers.MAX_OCTETS + 1)),
]:
    try:
        fichiers.lire_fichier(nom, b64(octets))
        check(label, False)
    except fichiers.ImportInvalide:
        check(label, True)

print("\n=== Import : doublons et ajout ===")
lignes = [
    {"nom": "Garage Lubin SARL", "telephone": "+590 690 11 22 33"},  # meme telephone que a1
    {"nom": "Jean M.", "email": "MME.JEAN@exemple.fr"},  # meme e-mail que a2
    {"nom": "garage lubin"},  # meme nom, sans contact
    {"nom": "Paul Durand", "telephone": "0690 44 55 66", "statut": "RDV", "relance": "15/11/2026", "montant": "abc"},
    {"nom": "Paul D.", "telephone": "06 90 44 55 66"},  # doublon dans le fichier
    {"nom": "", "telephone": "", "email": ""},
    {"nom": "", "email": "contact@hotel.fr"},
]
tri = fichiers.trier(ANCIEN, lignes, place=10)
check("3 doublons du carnet + 1 du fichier écartés", len(tri["doublons"]) == 4)
check("ligne vide ignorée", tri["ignores"] == 1)
check("2 à ajouter, le nom manquant prend l'e-mail", [p["nom"] for p in tri["a_ajouter"]] == ["Paul Durand", "contact@hotel.fr"])
paul = tri["a_ajouter"][0]
check("statut du fichier gardé, date française lue, montant illisible ignoré", paul["statut"] == "RDV" and paul["relance"] == "2026-11-15" and paul["montant"] is None)
check("plus de place : le reste est compté à part", fichiers.trier(ANCIEN, lignes, place=1)["hors_place"] == 1)

faux = FauxNotion([dict(x) for x in ANCIEN])
patches = avec(faux)
try:
    fichier = b64("Nom;Téléphone\nPaul Durand;0690 44 55 66\nGarage Lubin;0690112233\n".encode())
    code, r = statut(portal_main.portal_prospects_import, portal_main.ProspectsImportRequest(nom_fichier="l.csv", contenu=fichier), authorization=BEARER)
    check("aperçu : 1 à ajouter, 1 doublon, rien d'écrit", code == 200 and r["apercu"]["a_ajouter"] == 1 and r["apercu"]["doublons"] == 1 and not r["confirme"] and faux.ecritures == 0)

    code, r = statut(portal_main.portal_prospects_import, portal_main.ProspectsImportRequest(nom_fichier="l.csv", contenu=fichier, confirmer=True), authorization=BEARER)
    check("confirmer : écrit une fois, nouveau en tête", code == 200 and r["confirme"] and faux.ecritures == 1 and r["carnet"]["prospects"][0]["nom"] == "Paul Durand")
    check("les prospects existants sont gardés", [p["id"] for p in faux.prospects()][1:] == ["a1", "a2"])
    nouveau = faux.prospects()[0]
    check("nouveau prospect : id, date d'ajout, À contacter", nouveau["id"] and nouveau["cree_le"] == AUJOURD_HUI and nouveau["statut"] == "À contacter")

    code, r = statut(portal_main.portal_prospects_import, portal_main.ProspectsImportRequest(nom_fichier="l.csv", contenu=fichier, confirmer=True), authorization=BEARER)
    check("réimporter le même fichier n'ajoute rien", code == 200 and r["apercu"]["a_ajouter"] == 0 and faux.ecritures == 1)

    code, detail = statut(portal_main.portal_prospects_import, portal_main.ProspectsImportRequest(nom_fichier="l.doc", contenu=fichier), authorization=BEARER)
    check("format refusé : 422 avec message clair", code == 422 and "Excel" in detail)

    reponse = portal_main.portal_prospects_export(format="xlsx", authorization=BEARER)
    check("export Excel : fichier en pièce jointe", reponse.media_type.endswith("spreadsheetml.sheet") and "attachment" in reponse.headers["content-disposition"])
    reponse = portal_main.portal_prospects_export(format="csv", authorization=BEARER)
    check("export CSV : 3 prospects", reponse.body.decode("utf-8-sig").count("\r\n") == 4)
    code, _ = statut(portal_main.portal_prospects_export, format="pdf", authorization=BEARER)
    check("export : format inconnu refusé", code == 422)
finally:
    fin(patches)

code, _ = statut(portal_main.portal_prospects_export, format="csv", authorization="")
check("export sans session : refusé", code == 401)
code, _ = statut(portal_main.portal_prospects_import, portal_main.ProspectsImportRequest(nom_fichier="l.csv", contenu="eA=="), authorization="Bearer faux")
check("import sans session valide : refusé", code == 401)

print("\n=== Historique des échanges ===")
faux = FauxNotion([dict(x) for x in ANCIEN])
patches = avec(faux)
try:
    v = carnet.vue_client(carnet.appliquer(CLIENT, "Marie Test", DASHBOARD, "noter", "a2", {"texte": "  Appelée, rappeler mardi ", "date": "2026-10-20", "relance": "2026-11-04"}))
    jean = next(x for x in v["prospects"] if x["id"] == "a2")
    check("échange noté, texte nettoyé, date gardée", jean["journal"][0]["texte"] == "Appelée, rappeler mardi" and jean["journal"][0]["date"] == "2026-10-20")
    check("la prochaine relance suit l'échange", jean["relance"] == "2026-11-04")
    v = carnet.vue_client(carnet.appliquer(CLIENT, "Marie Test", DASHBOARD, "noter", "a2", {"texte": "Devis envoyé", "date": "2026-10-10"}))
    jean = next(x for x in v["prospects"] if x["id"] == "a2")
    check("historique trié par date", [e["texte"] for e in jean["journal"]] == ["Devis envoyé", "Appelée, rappeler mardi"])
    v = carnet.vue_client(carnet.appliquer(CLIENT, "Marie Test", DASHBOARD, "oublier", "a2", {"echange_id": jean["journal"][0]["id"]}))
    check("effacer un échange", [e["texte"] for e in next(x for x in v["prospects"] if x["id"] == "a2")["journal"]] == ["Appelée, rappeler mardi"])

    for label, action, donnees, erreur in [
        ("échange vide refusé", "noter", {"texte": "  "}, carnet.CarnetInvalide),
        ("échange daté dans le futur refusé", "noter", {"texte": "x", "date": "2099-01-01"}, carnet.CarnetInvalide),
        ("échange inconnu : LookupError", "oublier", {"echange_id": "zzz"}, LookupError),
    ]:
        try:
            carnet.appliquer(CLIENT, "Marie Test", DASHBOARD, action, "a2", donnees)
            check(label, False)
        except erreur:
            check(label, True)

    lignes = list(csv.reader(io.StringIO(fichiers.export_csv(faux.prospects()).decode("utf-8-sig")), delimiter=";"))
    check("export : colonne historique", lignes[0][-1] == "Historique des échanges" and "Appelée, rappeler mardi" in lignes[2][-1])
finally:
    fin(patches)

print("\n=== Entonnoir ===")
e = carnet.entonnoir([
    {"statut": "À contacter"}, {"statut": "À contacter"},
    {"statut": "Contacté"}, {"statut": "Contacté"},
    {"statut": "RDV"},
    {"statut": "À contacter", "evenements": [{"statut": "RDV"}]},  # revenu en arrière : compte quand même
    {"statut": "Client", "evenements": []},  # client direct
    {"statut": "Client", "evenements": [{"statut": "Contacté"}, {"statut": "RDV"}, {"statut": "Client"}]},
])
check("comptes par étape atteinte", (e["prospects"], e["contactes"], e["rdv"], e["clients"]) == (8, 6, 4, 2))
check("taux", (e["taux_contact"], e["taux_rdv"], e["taux_signature"], e["taux_global"]) == (75, 67, 50, 33))
check("carnet vide : pas de taux", carnet.entonnoir([])["taux_rdv"] is None)
ligne = carnet.resume_pour_idris([{"statut": "Client", "montant": 900, "evenements": []}, {"statut": "Contacté"}])
check("ligne pour Idris", ligne == "Son carnet de prospects : 2 prospects, 2 contactés, 1 RDV obtenus, 1 clients signés (taux contact → RDV 50 %, RDV → client 100 %), 900 € signés.")
check("pas de ligne si carnet vide", carnet.resume_pour_idris([]) is None)
check("vue_client donne l'entonnoir", carnet.vue_client({"prospects": ANCIEN})["entonnoir"]["clients"] == 1)

print("\n=== Export côté coach ===")
faux = FauxNotion([dict(x) for x in ANCIEN])
patches = avec(faux)
try:
    code, _ = statut(portal_main.coach_eleve_prospects_export, CLIENT, format="csv", x_coach_key="mauvaise")
    check("sans la clé coach : 401", code == 401)
    reponse = portal_main.coach_eleve_prospects_export(CLIENT, format="xlsx", x_coach_key="cle-coach-crm")
    check("avec la clé : fichier Excel", reponse.media_type.endswith("spreadsheetml.sheet") and load_workbook(io.BytesIO(reponse.body)).active["A2"].value == "Garage Lubin")
    code, _ = statut(portal_main.coach_eleve_prospects_export, CLIENT, format="pdf", x_coach_key="cle-coach-crm")
    check("format inconnu : 422", code == 422)
finally:
    fin(patches)

print(f"\n{passed} PASS, {failed} FAIL")
sys.exit(1 if failed else 0)
