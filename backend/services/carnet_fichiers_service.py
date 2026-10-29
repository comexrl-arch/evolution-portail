# Export et import du carnet de prospects (demande de Rony du 29/10/2026 :
# « pour un CRM il y a l'exportation »). L'eleve telecharge son carnet en
# Excel ou en CSV, et peut y verser une liste qu'il a deja (Excel, CSV, ou
# contacts exportes de son telephone en .vcf). Les doublons sont ecartes :
# meme telephone, meme e-mail, ou meme nom quand la ligne n'a ni l'un ni
# l'autre.

import base64
import csv
import io
import re
import unicodedata
from datetime import date, datetime

from backend.services import prospects_service as carnet

MAX_OCTETS = 2 * 1024 * 1024
MAX_LIGNES = 2000

COLONNES_EXPORT = [
    ("Nom", "nom"),
    ("Téléphone", "telephone"),
    ("E-mail", "email"),
    ("Statut", "statut"),
    ("Prochaine relance", "relance"),
    ("Source", "source"),
    ("Montant (€)", "montant"),
    ("Note", "note"),
    ("Ajouté le", "cree_le"),
    ("Dernier contact", "dernier_contact"),
    ("Date du RDV", "date_rdv"),
    ("Devenu client le", "date_client"),
]

# En-tetes reconnus a l'import, une fois sans accents ni ponctuation. Couvre
# l'export du portail, les tableurs faits main et l'export Google Contacts.
_ALIAS = {
    "nom": {"nom", "name", "nomduprospect", "prospect", "contact", "nomcomplet", "fullname", "entreprise", "societe", "raisonsociale", "company", "organisation", "organizationname", "organization1name"},
    "prenom": {"prenom", "firstname", "givenname"},
    "nom_famille": {"nomdefamille", "lastname", "familyname", "surname"},
    "telephone": {"telephone", "tel", "phone", "mobile", "portable", "gsm", "numero", "numerodetelephone", "phonenumber", "phone1value", "telephonemobile"},
    "email": {"email", "mail", "courriel", "adresseemail", "adressemail", "emailaddress", "email1value"},
    "statut": {"statut", "status", "etape", "stage", "statutcontact"},
    "relance": {"relance", "prochainerelance", "daterelance", "followup", "prochaincontact"},
    "source": {"source", "origine", "canal", "provenance"},
    "montant": {"montant", "montanteur", "valeur", "valeurconversion", "ca", "chiffredaffaires", "amount", "prix"},
    "note": {"note", "notes", "commentaire", "commentaires", "remarque", "remarques", "comment", "prochaineaction"},
}


class ImportInvalide(ValueError):
    def __init__(self, message_client: str):
        super().__init__(message_client)
        self.message_client = message_client


def _cle(texte) -> str:
    sans_accents = unicodedata.normalize("NFKD", str(texte or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", sans_accents.lower())


# --- Export ----------------------------------------------------------------------

def _date_statut(prospect: dict, statut: str) -> str:
    dates = [e.get("date") for e in prospect.get("evenements") or [] if e.get("statut") == statut and e.get("date")]
    return max(dates) if dates else ""


def lignes_export(prospects: list[dict]) -> list[list]:
    lignes = []

    for brut in prospects:
        p = carnet.completer(brut)
        valeurs = {
            **p,
            "relance": p.get("relance") or "",
            "montant": p.get("montant") if p.get("montant") is not None else "",
            "dernier_contact": _date_statut(p, "Contacté"),
            "date_rdv": _date_statut(p, "RDV"),
            "date_client": _date_statut(p, "Client"),
        }
        lignes.append(["" if valeurs.get(cle) is None else valeurs[cle] for _, cle in COLONNES_EXPORT])

    return lignes


def _anti_formule(valeur):
    # Un tableur execute une cellule qui commence par = + - @ : on la neutralise.
    if isinstance(valeur, str) and valeur[:1] in ("=", "+", "-", "@") and not re.fullmatch(r"\+?[\d .()-]+", valeur):
        return "'" + valeur
    return valeur


def export_csv(prospects: list[dict]) -> bytes:
    sortie = io.StringIO()
    # Point-virgule et BOM : Excel en francais l'ouvre directement en colonnes.
    ecrivain = csv.writer(sortie, delimiter=";", lineterminator="\r\n")
    ecrivain.writerow([titre for titre, _ in COLONNES_EXPORT])

    for ligne in lignes_export(prospects):
        ecrivain.writerow([str(_anti_formule(v)).replace(".", ",") if isinstance(v, float) else _anti_formule(v) for v in ligne])

    return ("﻿" + sortie.getvalue()).encode("utf-8")


def export_xlsx(prospects: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Mes prospects"
    feuille.append([titre for titre, _ in COLONNES_EXPORT])

    for cellule in feuille[1]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor="1F2A44")

    colonnes_dates = {i for i, (_, cle) in enumerate(COLONNES_EXPORT) if cle in ("relance", "cree_le", "dernier_contact", "date_rdv", "date_client")}

    for ligne in lignes_export(prospects):
        valeurs = []

        for i, valeur in enumerate(ligne):
            if i in colonnes_dates and valeur:
                try:
                    valeur = date.fromisoformat(str(valeur)[:10])
                except ValueError:
                    pass
            valeurs.append(_anti_formule(valeur))
        feuille.append(valeurs)

    for i, (titre, cle) in enumerate(COLONNES_EXPORT, start=1):
        lettre = feuille.cell(row=1, column=i).column_letter
        feuille.column_dimensions[lettre].width = {"nom": 28, "note": 45, "email": 28}.get(cle, max(14, len(titre) + 2))

        if i - 1 in colonnes_dates:
            for (cellule,) in feuille.iter_rows(min_row=2, min_col=i, max_col=i):
                cellule.number_format = "DD/MM/YYYY"

    feuille.freeze_panes = "A2"
    feuille.auto_filter.ref = feuille.dimensions
    sortie = io.BytesIO()
    classeur.save(sortie)
    return sortie.getvalue()


def nom_fichier(nom_client: str, extension: str) -> str:
    base = _cle(nom_client)[:30] or "eleve"
    return f"carnet-prospects-{base}-{date.today().isoformat()}.{extension}"


# --- Import : lecture du fichier ---------------------------------------------------

def _texte_fichier(octets: bytes) -> str:
    for encodage in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return octets.decode(encodage)
        except UnicodeDecodeError:
            continue
    raise ImportInvalide("Je n'arrive pas à lire ce fichier.")


def _lire_csv(octets: bytes) -> list[list[str]]:
    texte = _texte_fichier(octets)
    premiere = texte.split("\n", 1)[0]
    separateur = max([";", ",", "\t"], key=premiere.count)
    return [ligne for ligne in csv.reader(io.StringIO(texte), delimiter=separateur)]


def _lire_xlsx(octets: bytes) -> list[list]:
    from openpyxl import load_workbook

    try:
        classeur = load_workbook(io.BytesIO(octets), read_only=True, data_only=True)
    except Exception:
        raise ImportInvalide("Je n'arrive pas à ouvrir ce fichier Excel. Enregistre-le en .xlsx ou en .csv et réessaie.")

    feuille = classeur.worksheets[0]
    lignes = []

    for ligne in feuille.iter_rows(values_only=True):
        lignes.append(list(ligne))

        if len(lignes) > MAX_LIGNES + 1:
            break

    classeur.close()
    return lignes


def _lire_vcf(octets: bytes) -> list[dict]:
    # Cartes de contact exportees d'un telephone : FN, ORG, TEL, EMAIL, NOTE.
    texte = re.sub(r"\r?\n[ \t]", "", _texte_fichier(octets))
    contacts = []
    courant = None

    for ligne in texte.splitlines():
        if ligne.upper().startswith("BEGIN:VCARD"):
            courant = {}
        elif ligne.upper().startswith("END:VCARD"):
            if courant is not None:
                contacts.append(courant)
            courant = None
        elif courant is not None and ":" in ligne:
            cle, valeur = ligne.split(":", 1)
            nom = cle.split(";", 1)[0].upper().split(".")[-1]
            valeur = valeur.replace("\\,", ",").replace("\\;", ";").replace("\\n", " ").strip()

            if nom == "FN":
                courant["nom"] = valeur
            elif nom == "N" and not courant.get("nom"):
                parties = valeur.split(";")
                courant["nom"] = " ".join(p for p in (parties[1:2] + parties[:1]) if p).strip()
            elif nom == "ORG":
                courant.setdefault("entreprise", valeur.replace(";", " ").strip())
            elif nom == "TEL":
                courant.setdefault("telephone", valeur)
            elif nom == "EMAIL":
                courant.setdefault("email", valeur)
            elif nom == "NOTE":
                courant.setdefault("note", valeur)

    for contact in contacts:
        if not contact.get("nom"):
            contact["nom"] = contact.get("entreprise") or ""

    return contacts


def _depuis_tableau(lignes: list[list]) -> list[dict]:
    lignes = [l for l in lignes if any(str(v or "").strip() for v in l)]

    if not lignes:
        return []

    entetes = [_cle(v) for v in lignes[0]]
    colonnes = {}

    for champ, alias in _ALIAS.items():
        for i, entete in enumerate(entetes):
            if entete in alias and champ not in colonnes and i not in colonnes.values():
                colonnes[champ] = i

    if not ({"nom", "prenom", "nom_famille", "telephone", "email"} & colonnes.keys()):
        raise ImportInvalide("Je ne trouve pas de colonne « Nom », « Téléphone » ou « E-mail » sur la première ligne de ton fichier.")

    resultat = []

    for ligne in lignes[1:]:
        def valeur(champ):
            i = colonnes.get(champ)
            return ligne[i] if i is not None and i < len(ligne) and ligne[i] is not None else ""

        # « Prénom » + « Nom » (ou First/Last name) donnent le nom complet.
        nom = " ".join(str(v).strip() for v in (valeur("prenom"), valeur("nom") or valeur("nom_famille")) if str(v).strip())

        resultat.append({
            "nom": nom,
            "telephone": valeur("telephone"),
            "email": valeur("email"),
            "statut": valeur("statut"),
            "relance": valeur("relance"),
            "source": valeur("source"),
            "montant": valeur("montant"),
            "note": valeur("note"),
        })

    return resultat


def lire_fichier(nom: str, contenu_base64: str) -> list[dict]:
    try:
        octets = base64.b64decode(contenu_base64 or "", validate=False)
    except (ValueError, TypeError):
        raise ImportInvalide("Le fichier n'a pas pu être envoyé. Réessaie.")

    if not octets:
        raise ImportInvalide("Ce fichier est vide.")

    if len(octets) > MAX_OCTETS:
        raise ImportInvalide("Ce fichier est trop lourd (2 Mo maximum).")

    extension = (nom or "").lower().rsplit(".", 1)[-1]

    if extension in ("xlsx", "xlsm"):
        lignes = _depuis_tableau(_lire_xlsx(octets))
    elif extension == "xls":
        raise ImportInvalide("Ce vieux format Excel (.xls) n'est pas lu. Enregistre-le en .xlsx ou en .csv et réessaie.")
    elif extension in ("vcf", "vcard"):
        lignes = _lire_vcf(octets)
    elif extension in ("csv", "txt", "tsv"):
        lignes = _depuis_tableau(_lire_csv(octets))
    else:
        raise ImportInvalide("Envoie un fichier Excel (.xlsx), CSV (.csv) ou des contacts de téléphone (.vcf).")

    if len(lignes) > MAX_LIGNES:
        raise ImportInvalide(f"Ce fichier a plus de {MAX_LIGNES} lignes. Garde les prospects que tu veux vraiment suivre.")

    return lignes


# --- Import : tri et ajout au carnet -----------------------------------------------

def _chiffres_tel(texte) -> str:
    chiffres = re.sub(r"\D", "", str(texte or ""))
    return chiffres[-9:] if len(chiffres) >= 9 else ""


def _date_import(valeur) -> str | None:
    if isinstance(valeur, datetime):
        return valeur.date().isoformat()

    if isinstance(valeur, date):
        return valeur.isoformat()

    texte = str(valeur or "").strip()

    if not texte:
        return None

    trouve = re.fullmatch(r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})", texte)

    if trouve:
        jour, mois, annee = (int(x) for x in trouve.groups())
        annee += 2000 if annee < 100 else 0

        try:
            return date(annee, mois, jour).isoformat()
        except ValueError:
            return None

    return carnet._date_iso(texte)


_STATUTS_IMPORT = {_cle(s): s for s in carnet.STATUTS}
_STATUTS_IMPORT.update({"acontacter": "À contacter", "nouveau": "À contacter", "contacte": "Contacté", "relance": "Contacté", "rdv": "RDV", "rendezvous": "RDV", "client": "Client", "gagne": "Client", "converti": "Client"})


def _nettoyer(ligne: dict) -> dict | None:
    telephone = carnet._texte(ligne.get("telephone"), carnet._MAX_CONTACT)
    email = carnet._texte(ligne.get("email"), carnet._MAX_CONTACT).replace(" ", "").lower()

    if email and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
        email = ""

    nom = carnet._texte(ligne.get("nom"), carnet._MAX_NOM) or telephone or email

    if not nom:
        return None

    try:
        montant = carnet.montant(ligne.get("montant"))
    except carnet.CarnetInvalide:
        montant = None

    return {
        "nom": nom,
        "telephone": telephone,
        "email": email,
        "contact": telephone or email,
        "statut": _STATUTS_IMPORT.get(_cle(ligne.get("statut")), "À contacter"),
        "relance": _date_import(ligne.get("relance")),
        "source": carnet._texte(ligne.get("source"), carnet._MAX_SOURCE),
        "montant": montant,
        "note": carnet._texte(ligne.get("note"), carnet._MAX_NOTE),
    }


def trier(existants: list[dict], lignes: list[dict], place: int) -> dict:
    """Ce que l'import ajouterait : {"a_ajouter", "doublons", "ignores", "hors_place"}."""
    tels, emails, noms = set(), set(), set()

    def retenir(p):
        if _chiffres_tel(p.get("telephone") or p.get("contact")):
            tels.add(_chiffres_tel(p.get("telephone") or p.get("contact")))
        if p.get("email") or "@" in (p.get("contact") or ""):
            emails.add((p.get("email") or p.get("contact")).lower())
        noms.add(_cle(p.get("nom")))

    for p in existants:
        retenir(carnet.completer(p))

    a_ajouter, doublons, ignores, hors_place = [], [], 0, 0

    for ligne in lignes:
        propre = _nettoyer(ligne)

        if not propre:
            ignores += 1
            continue

        tel = _chiffres_tel(propre["telephone"])
        double = (tel and tel in tels) or (propre["email"] and propre["email"] in emails) or (
            not tel and not propre["email"] and _cle(propre["nom"]) in noms
        )

        if double:
            doublons.append(propre["nom"])
            continue

        if len(a_ajouter) >= place:
            hors_place += 1
            continue

        a_ajouter.append(propre)
        retenir(propre)

    return {"a_ajouter": a_ajouter, "doublons": doublons, "ignores": ignores, "hors_place": hors_place}


def resume(tri: dict) -> dict:
    return {
        "a_ajouter": len(tri["a_ajouter"]),
        "exemples": [p["nom"] for p in tri["a_ajouter"][:5]],
        "doublons": len(tri["doublons"]),
        "exemples_doublons": tri["doublons"][:5],
        "ignores": tri["ignores"],
        "hors_place": tri["hors_place"],
    }


def importer(client_page_id: str, nom_client: str, dashboard: dict, nom: str, contenu_base64: str, confirmer: bool) -> dict:
    """Sans confirmer : l'apercu de l'import. Avec : les prospects sont ajoutes
    en tete du carnet, dans l'ordre du fichier."""
    lignes = lire_fichier(nom, contenu_base64)
    existant = carnet.lire(client_page_id, dashboard)
    place = max(0, carnet.MAX_PROSPECTS - len(existant["prospects"]))
    tri = trier(existant["prospects"], lignes, place)
    reponse = {"apercu": resume(tri), "confirme": False}

    if confirmer and tri["a_ajouter"]:
        aujourd_hui = carnet._aujourd_hui().isoformat()
        nouveaux = [carnet.nouveau_prospect(p, aujourd_hui) for p in tri["a_ajouter"]]
        enregistre = carnet.enregistrer_liste(client_page_id, nom_client, existant, nouveaux + [dict(p) for p in existant["prospects"]])
        reponse.update({"confirme": True, "carnet": carnet.vue_client(enregistre)})

    return reponse
